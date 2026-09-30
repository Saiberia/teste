"""Human-readable output: Markdown (files, web), Telegram HTML, DOCX."""

from __future__ import annotations

import html
import io
import re

from .i18n import t
from .models import Answer, AnswerStatus, Item, ItemKind, ItemOrigin, MeetingReport
from .transcript import format_clock

TELEGRAM_LIMIT = 4000  # Telegram hard limit is 4096; keep a margin


def _item_label(item: Item, lang: str) -> str:
    kind = t("question" if item.kind == ItemKind.QUESTION else "task", lang)
    if item.origin == ItemOrigin.USER:
        return f"{kind} ({t('typed', lang)})"
    parts = [kind]
    if item.origin == ItemOrigin.VOICE:
        parts[0] += f" ({t('by_voice', lang)})"
    elif item.detector == "heuristic":
        parts[0] += f" ({t('possible', lang)})"
    if item.speaker:
        parts.append(item.speaker)
    if item.start is not None:
        parts.append(format_clock(item.start))
    return ", ".join(parts)


def _status_note(item: Item, answer: Answer | None, lang: str) -> str:
    if answer is None:
        return t("preparing" if item.is_command else "not_requested", lang)
    if answer.status == AnswerStatus.FAILED:
        return t("failed", lang) + (f": {answer.summary}" if answer.summary else "")
    if answer.status == AnswerStatus.NEEDS_LLM:
        return t("offline", lang)
    return f"{t('confidence', lang)}: {t(answer.confidence, lang)}; {t('check_before_sending', lang)}"


def ordered_items(report: MeetingReport) -> tuple[list[Item], list[Item]]:
    """User's own tasks first (the point of the product), then what was heard."""
    mine = [i for i in report.items if i.is_command]
    heard = [i for i in report.items if not i.is_command]
    return mine, heard


def _md_link(title: str, url: str) -> str:
    safe = re.sub(r"([\[\]\\])", r"\\\1", title)
    return f"[{safe}](<{url}>)"


def report_to_markdown(report: MeetingReport, lang: str = "ru") -> str:
    out = [f"# {report.title}", ""]
    recap = report.recap
    out += [f"## {t('recap', lang)}", recap.summary or "—", ""]
    if recap.decisions:
        out += [f"## {t('decisions', lang)}", *[f"- {d}" for d in recap.decisions], ""]
    if recap.action_items:
        out.append(f"## {t('action_items', lang)}")
        for a in recap.action_items:
            meta = ", ".join(x for x in (a.owner, a.due) if x)
            out.append(f"- {a.text}" + (f" ({meta})" if meta else ""))
        out.append("")
    for section in recap.sections:
        if section.bullets:
            out += [f"## {section.title}", *[f"- {b}" for b in section.bullets], ""]
    mine, heard = ordered_items(report)
    n = 0
    for heading, items in ((t("my_tasks", lang), mine), (t("heard", lang), heard)):
        out.append(f"## {heading} ({len(items)})")
        if not items:
            out += [t("nothing_found", lang), ""]
        for item in items:
            n += 1
            ans = report.answer_for(item.id)
            out += ["", f"### {n}. {item.text}", f"*{_item_label(item, lang)}*"]
            if item.quote and item.quote != item.text:
                out.append(f"> {item.quote}")
            out += ["", f"_{_status_note(item, ans, lang)}_"]
            if ans and ans.status != AnswerStatus.FAILED:
                if ans.summary:
                    out += ["", f"**{t('short', lang)}:** {ans.summary}"]
                if ans.body:
                    out += ["", ans.body]
                if ans.assumptions:
                    out += ["", f"**{t('assumptions', lang)}:**", *[f"- {x}" for x in ans.assumptions]]
                if ans.warnings:
                    out += ["", f"**{t('ai_check', lang)}:**", *[f"- ⚠ {x}" for x in ans.warnings]]
                if ans.sources:
                    out += ["", f"**{t('sources', lang)}:**"]
                    for s in ans.sources:
                        out.append(f"- {_md_link(s.title, s.ref)}" if s.ref.startswith("http") else f"- {s.title} ({s.ref})")
        out.append("")
    out += ["---", f"_{t('mode', lang)}: {report.mode}. {t('footer', lang)}_"]
    return "\n".join(out).strip() + "\n"


# --- Telegram -------------------------------------------------------------------

def _h(text: str) -> str:
    return html.escape(text, quote=False)


_TAG_RE = re.compile(r"</?([a-z]+)(?:\s[^>]*)?>")


def _balanced(fragment: str) -> bool:
    stack: list[str] = []
    for match in _TAG_RE.finditer(fragment):
        tag = match.group(1)
        if match.group(0).startswith("</"):
            if not stack or stack.pop() != tag:
                return False
        else:
            stack.append(tag)
    return not stack


def _strip_tags(fragment: str) -> str:
    return _TAG_RE.sub("", fragment)


def _safe_cut(line: str, limit: int) -> int:
    """Cut position <= limit that doesn't split an HTML tag or an &entity;."""
    cut = limit
    lt, gt = line.rfind("<", 0, cut), line.rfind(">", 0, cut)
    if lt > gt:
        cut = lt
    amp = line.rfind("&", max(0, cut - 8), cut)
    if amp != -1 and ";" not in line[amp:cut]:
        cut = amp
    return cut if cut > 0 else limit


def _hard_split(block: str, limit: int) -> list[str]:
    out, buf = [], ""
    for line in block.split("\n"):
        while len(line) > limit:
            if buf:
                out.append(buf)
                buf = ""
            cut = _safe_cut(line, limit)
            out.append(line[:cut])
            line = line[cut:]
        candidate = f"{buf}\n{line}" if buf else line
        if len(candidate) <= limit:
            buf = candidate
        else:
            out.append(buf)
            buf = line
    if buf:
        out.append(buf)
    return out


def split_message(text: str, limit: int = TELEGRAM_LIMIT) -> list[str]:
    """Split on blank lines, then lines, then safe hard cuts. Parts whose tags got
    unbalanced by a cut are sent without markup rather than rejected by Telegram."""
    parts: list[str] = []
    buf = ""
    for block in text.split("\n\n"):
        pieces = [block] if len(block) <= limit else _hard_split(block, limit)
        for piece in pieces:
            candidate = f"{buf}\n\n{piece}" if buf else piece
            if len(candidate) <= limit:
                buf = candidate
            else:
                if buf:
                    parts.append(buf)
                buf = piece
    if buf:
        parts.append(buf)
    return [p if _balanced(p) else _strip_tags(p) for p in parts]


def item_to_telegram(n: int, item: Item, ans: Answer | None, lang: str = "ru") -> str:
    lines = [f"<b>{n}. {_h(item.text)}</b>", f"<i>{_h(_item_label(item, lang))}</i>",
             f"<i>{_h(_status_note(item, ans, lang))}</i>"]
    if ans is None or ans.status == AnswerStatus.FAILED:
        return "\n".join(lines)
    if ans.summary:
        lines += ["", f"<b>{t('short', lang)}:</b> {_h(ans.summary)}"]
    if ans.body:
        lines += ["", _h(ans.body)]
    if ans.assumptions:
        lines += ["", f"<b>{t('assumptions', lang)}:</b>", *[f"• {_h(x)}" for x in ans.assumptions]]
    if ans.warnings:
        lines += ["", f"<b>{t('ai_check', lang)}:</b>", *[f"⚠ {_h(x)}" for x in ans.warnings]]
    if ans.sources:
        lines += ["", f"<b>{t('sources', lang)}:</b>"]
        for s in ans.sources:
            if s.ref.startswith("http"):
                lines.append(f'• <a href="{html.escape(s.ref, quote=True)}">{_h(s.title)}</a>')
            else:
                lines.append(f"• {_h(s.title)}")
    return "\n".join(lines)


def report_to_telegram(report: MeetingReport, lang: str = "ru") -> list[str]:
    """Telegram HTML messages: summary first, then the user's tasks, then suggestions."""
    r = report.recap
    head = [f"<b>{_h(report.title)}</b>", "", f"<b>{t('recap', lang)}</b>", _h(r.summary or "—")]
    if r.decisions:
        head += ["", f"<b>{t('decisions', lang)}</b>", *[f"• {_h(d)}" for d in r.decisions]]
    if r.action_items:
        head += ["", f"<b>{t('action_items', lang)}</b>"]
        for a in r.action_items:
            meta = ", ".join(x for x in (a.owner, a.due) if x)
            head.append(f"• {_h(a.text)}" + (f" <i>({_h(meta)})</i>" if meta else ""))
    for section in r.sections:
        if section.bullets:
            head += ["", f"<b>{_h(section.title)}</b>", *[f"• {_h(b)}" for b in section.bullets]]
    mine, heard = ordered_items(report)
    head += ["", f"{t('my_tasks', lang)}: {len(mine)} · {t('heard', lang)}: {len(heard)}"]
    blocks = ["\n".join(head)]
    for n, item in enumerate(mine + heard, 1):
        blocks.append(item_to_telegram(n, item, report.answer_for(item.id), lang))
    messages: list[str] = []
    for block in blocks:
        messages.extend(split_message(block))
    return messages


# --- DOCX (like mymeet.ai export) ---------------------------------------------------

def report_to_docx(report: MeetingReport, lang: str = "ru") -> bytes:
    """Word export built from the Markdown report (headings, bullets, quotes, text)."""
    from docx import Document

    doc = Document()
    for line in report_to_markdown(report, lang).splitlines():
        stripped = line.strip()
        if not stripped or stripped == "---":
            continue
        level = len(stripped) - len(stripped.lstrip("#"))
        if 1 <= level <= 3 and stripped[level:level + 1] == " ":
            doc.add_heading(stripped[level + 1:], level=level)
        elif stripped.startswith(("- ", "• ")):
            doc.add_paragraph(_plain(stripped[2:]), style="List Bullet")
        elif stripped.startswith("> "):
            doc.add_paragraph(_plain(stripped[2:]), style="Quote")
        else:
            doc.add_paragraph(_plain(stripped))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plain(md: str) -> str:
    md = re.sub(r"\[([^\]]*)\]\(<?([^)>]*)>?\)", r"\1 (\2)", md)
    return md.replace("**", "").replace("__", "").strip("_* ")


def markdown_to_html(md: str) -> str:
    """Tiny safe Markdown: escaped text, paragraphs, bullet lists, **bold**, headings."""
    out, in_list = [], False
    for line in (md or "").splitlines():
        t = html.escape(line.strip())
        t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", t)
        bullet = re.match(r"^[-*•]\s+(.*)$", t) or re.match(r"^\d+[.)]\s+(.*)$", t)
        if bullet:
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{bullet.group(1)}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        head = re.match(r"^#{1,4}\s+(.*)$", t)
        if head:
            out.append(f"<h4>{head.group(1)}</h4>")
        elif t:
            out.append(f"<p>{t}</p>")
    if in_list:
        out.append("</ul>")
    return "".join(out)


_PALETTE = ["#0f8b7d", "#7c5cd6", "#d9730d", "#2f6fdb", "#c2417a", "#4d8b2f", "#8a6d1f"]


def report_to_html(report: MeetingReport, lang: str = "ru") -> str:
    """Standalone, printable page (Print -> Save as PDF): recap, tasks with answers, full transcript."""
    e = html.escape
    r = report.recap
    answers = {a.item_id: a for a in report.answers}
    colors: dict[str, str] = {}
    for seg in report.segments:
        colors.setdefault(seg.speaker or "—", _PALETTE[len(colors) % len(_PALETTE)])
    mine, heard = ordered_items(report)
    parts = [f"<h1>{e(report.title)}</h1><p class='meta'>{e(str(report.created_at)[:16].replace('T', ' '))}"
             f" · {len(report.segments)} реплик · участники: "
             + ", ".join(f"<span class='who' style='--c:{c}'>{e(s)}</span>" for s, c in colors.items()) + "</p>"]
    if r and r.summary:
        parts.append(f"<section><h2>Итоги</h2><p>{e(r.summary)}</p>")
        if r.decisions:
            parts.append("<h3>Решения</h3><ul>" + "".join(f"<li>{e(d)}</li>" for d in r.decisions) + "</ul>")
        if r.action_items:
            parts.append("<h3>Поручения</h3><ul>" + "".join(
                f"<li>{e(a.text)}" + (f" <span class='muted'>— {e(', '.join(x for x in (a.owner, a.due) if x))}</span>"
                                      if a.owner or a.due else "") + "</li>" for a in r.action_items) + "</ul>")
        for sec in r.sections or []:
            if sec.bullets:
                parts.append(f"<h3>{e(sec.title)}</h3><ul>" + "".join(f"<li>{e(b)}</li>" for b in sec.bullets) + "</ul>")
        parts.append("</section>")
    items = mine + heard
    if items:
        parts.append("<section><h2>Задачи и вопросы</h2>")
        for it in items:
            ans = answers.get(it.id)
            body = markdown_to_html(ans.text) if ans and ans.text else ""
            who = f"{e(it.speaker)} · " if it.speaker else ""
            parts.append(f"<div class='card'><div class='q'>{e(it.text)}</div><div class='muted small'>{who}"
                         f"{format_clock(it.start) if it.start is not None else ''}</div>"
                         + (f"<div class='ans'>{body}</div>" if body else "") + "</div>")
        parts.append("</section>")
    parts.append("<section><h2>Расшифровка</h2><div class='tr'>")
    for seg in report.segments:
        spk = seg.speaker or "—"
        parts.append(f"<div class='line'><span class='t'>{format_clock(seg.start) if seg.start is not None else ''}</span>"
                     f"<span class='who' style='--c:{colors[spk]}'>{e(spk)}</span><span class='x'>{e(seg.text)}</span></div>")
    parts.append("</div></section>")
    css = """
:root{--bg:#f6f7f9;--card:#fff;--ink:#15191f;--muted:#6b7380;--line:#e4e7ec}
@media (prefers-color-scheme:dark){:root{--bg:#111418;--card:#1a1e24;--ink:#e8ebef;--muted:#9aa3ae;--line:#2a3038}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.6 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:860px;margin:0 auto;padding:40px 20px}h1{font-size:28px;margin:0 0 6px}h2{font-size:18px;margin:0 0 12px}
h3{font-size:15px;margin:16px 0 6px}section{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:20px 22px;margin:18px 0}
.meta,.muted{color:var(--muted)}.small{font-size:13px}.who{color:var(--c);font-weight:600}
.card{border-left:3px solid #0f8b7d;padding:8px 14px;margin:12px 0}.q{font-weight:600}.ans{margin-top:6px}
.line{display:grid;grid-template-columns:52px 150px 1fr;gap:10px;padding:5px 0;border-bottom:1px solid var(--line)}
.t{color:var(--muted);font:12px/1.9 ui-monospace,Consolas,monospace}
@media (max-width:600px){.line{grid-template-columns:44px 1fr}.line .x{grid-column:1/-1}}
@media print{body{background:#fff;color:#000}section{break-inside:auto;border-color:#ccc}.card{break-inside:avoid}}
"""
    return (f"<!doctype html><html lang='{lang}'><head><meta charset='utf-8'><meta name='viewport' "
            f"content='width=device-width,initial-scale=1'><title>{e(report.title)}</title><style>{css}</style></head>"
            f"<body><main>{''.join(parts)}<p class='muted small'>Recapper · чтобы сохранить в PDF: Ctrl+P → «Сохранить как PDF»</p>"
            f"</main></body></html>")
